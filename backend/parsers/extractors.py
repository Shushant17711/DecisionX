"""
DecisionX — Document text extractors (pure, synchronous)

`parse_bytes` turns one uploaded file into plain text. It imports nothing from
the web stack so it can run inside a throwaway worker process (see worker.py),
and every format is read with explicit ceilings so a hostile file costs a
bounded amount of CPU and memory:

  * zip containers (xlsx/docx) are inspected before parsing: entry count, total
    uncompressed size, per-member size and compression ratio (zip bombs).
  * xlsx is streamed with openpyxl read_only — at most MAX_SHEETS sheets and
    MAX_SCAN_ROWS rows each are ever read; pandas is not used.
  * PDFs are opened with only the first MAX_PDF_PAGES pages.
  * text stops at `max_chars`.

Returns (text, error, truncated) where `truncated` means the source held more
than was extracted.
"""

import csv
import io
import math
import zipfile
from xml.etree import ElementTree

MAX_FILE_BYTES = 10 * 1024 * 1024

MAX_ZIP_ENTRIES = 2000
MAX_ZIP_TOTAL_UNCOMPRESSED = 200 * 1024 * 1024
MAX_ZIP_MEMBER_UNCOMPRESSED = 64 * 1024 * 1024
MAX_ZIP_RATIO = 150
_RATIO_CHECK_FLOOR = 1024 * 1024

MAX_SHEETS = 5
MAX_SCAN_ROWS = 2000
MAX_SHOWN_ROWS = 25
MAX_COLS = 40
MAX_CELL_CHARS = 80

MAX_PDF_PAGES = 40

TEXT_SUFFIXES = (".txt", ".md", ".json", ".yaml", ".yml", ".tsv", ".log", ".xml", ".html")
SUPPORTED_SUFFIXES = (".pdf", ".csv", ".xlsx", ".docx") + TEXT_SUFFIXES


class Rejected(Exception):
    """A file we refuse to parse, with a message fit for the user."""


def parse_bytes(filename: str, content: bytes, max_chars: int) -> tuple[str, str | None, bool]:
    if not content:
        return "", "File is empty", False
    if len(content) > MAX_FILE_BYTES:
        return "", f"File exceeds the {MAX_FILE_BYTES // (1024 * 1024)}MB limit", False

    name = filename.lower()
    try:
        if name.endswith(".pdf"):
            text, truncated = _extract_pdf(content, max_chars)
        elif name.endswith(".csv"):
            text, truncated = _extract_csv(content, max_chars)
        elif name.endswith(".xlsx"):
            text, truncated = _extract_xlsx(content, max_chars)
        elif name.endswith(".docx"):
            text, truncated = _extract_docx(content, max_chars)
        elif name.endswith(".xls"):
            return "", "Legacy .xls files are not supported — save as .xlsx or CSV and re-upload", False
        elif name.endswith((".doc", ".ppt", ".pptx")):
            return "", "This Office format is not supported — upload PDF, DOCX, XLSX, CSV, or text", False
        elif name.endswith(TEXT_SUFFIXES):
            if _looks_binary(content):
                return "", "This file is not readable text (binary content with a text extension)", False
            text = decode_text(content)
            text, truncated = text[:max_chars], len(text) > max_chars
        else:
            if _looks_binary(content):
                return "", "Unsupported file type — upload a PDF, CSV, spreadsheet, or text file", False
            text = decode_text(content)
            text, truncated = text[:max_chars], len(text) > max_chars
    except Rejected as e:
        return "", str(e), False
    except Exception as e:  # noqa: BLE001 — any parser failure is a per-file error
        return "", f"Could not parse: {type(e).__name__}", False

    text = text.strip()
    if not text:
        return "", "No extractable text", False
    return text, None, truncated


# ── Text decoding ───────────────────────────────────────────────────────────

def decode_text(content: bytes) -> str:
    """BOM-aware decode: UTF-8(-sig), UTF-16/32 with BOM, then cp1252, then latin-1."""
    if content.startswith(b"\xef\xbb\xbf"):
        return content[3:].decode("utf-8", errors="replace")
    if content.startswith((b"\xff\xfe\x00\x00", b"\x00\x00\xfe\xff")):
        return content.decode("utf-32", errors="replace")
    if content.startswith((b"\xff\xfe", b"\xfe\xff")):
        return content.decode("utf-16", errors="replace")
    for encoding in ("utf-8", "cp1252"):
        try:
            return content.decode(encoding)
        except UnicodeDecodeError:
            continue
    return content.decode("latin-1")


def _looks_binary(content: bytes) -> bool:
    sample = content[:8192]
    if sample.startswith((b"\xff\xfe", b"\xfe\xff", b"\xef\xbb\xbf")):
        return False
    if b"\x00" in sample:
        return True
    # Real binaries are full of NULs and control bytes. Undecodable high bytes
    # alone are not a signal: that is exactly what cp1252/latin-1 text looks like.
    control = sum(1 for b in sample if b < 32 and b not in (9, 10, 12, 13))
    return control > len(sample) * 0.05


# ── Zip containers ──────────────────────────────────────────────────────────

def _open_zip(content: bytes, kind: str) -> zipfile.ZipFile:
    """Open an OOXML container after checking it cannot expand into a bomb."""
    try:
        archive = zipfile.ZipFile(io.BytesIO(content))
    except zipfile.BadZipFile:
        raise Rejected(f"Not a valid {kind} file (corrupt or mislabelled)") from None

    infos = archive.infolist()
    if len(infos) > MAX_ZIP_ENTRIES:
        raise Rejected(f"{kind} file has too many internal parts ({len(infos)})")
    total = 0
    for info in infos:
        total += info.file_size
        if info.file_size > MAX_ZIP_MEMBER_UNCOMPRESSED:
            raise Rejected(f"{kind} file expands too large to read safely")
        if info.file_size > _RATIO_CHECK_FLOOR and info.file_size > MAX_ZIP_RATIO * max(1, info.compress_size):
            raise Rejected(f"{kind} file is suspiciously compressed (possible zip bomb)")
    if total > MAX_ZIP_TOTAL_UNCOMPRESSED:
        raise Rejected(f"{kind} file expands too large to read safely")
    return archive


# ── Tables (CSV / XLSX) ─────────────────────────────────────────────────────

def _cell(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return " ".join(str(value).split())[:MAX_CELL_CHARS]


def _as_number(value) -> float | None:
    if isinstance(value, bool) or value is None:
        return None
    if isinstance(value, (int, float)):
        number = float(value)
    else:
        try:
            number = float(str(value).replace(",", "").strip())
        except ValueError:
            return None
    return number if math.isfinite(number) else None


def _summarise_table(label: str, rows_iter, approx_rows: int | None) -> tuple[str, bool]:
    """Header, first rows, and per-column numeric stats from at most MAX_SCAN_ROWS rows."""
    header: list[str] | None = None
    shown: list[str] = []
    stats: dict[int, list[float]] = {}
    scanned = 0
    more = False

    for row in rows_iter:
        if scanned >= MAX_SCAN_ROWS:
            more = True
            break
        cells = list(row)[:MAX_COLS]
        if header is None:
            if not any(c not in (None, "") for c in cells):
                continue
            header = [_cell(c) or f"col{i + 1}" for i, c in enumerate(cells)]
            continue
        scanned += 1
        if len(shown) < MAX_SHOWN_ROWS:
            shown.append("\t".join(_cell(c) for c in cells))
        for i, c in enumerate(cells):
            number = _as_number(c)
            if number is not None:
                stats.setdefault(i, []).append(number)

    if header is None:
        return "", False

    count_note = f"{scanned}{'+' if more else ''} data rows"
    if approx_rows and approx_rows - 1 > scanned:
        count_note = f"about {approx_rows - 1} data rows (first {scanned} read)"
    parts = [
        f"{label}: {count_note}, {len(header)} columns.",
        "Columns: " + ", ".join(header),
        "",
        f"First {len(shown)} rows:",
        "\t".join(header),
        *shown,
    ]
    numeric = [
        (header[i] if i < len(header) else f"col{i + 1}", vals)
        for i, vals in sorted(stats.items())
        if len(vals) >= max(1, scanned // 2)
    ]
    if numeric:
        parts += ["", "Numeric summary (rows read):"]
        for name, vals in numeric:
            parts.append(
                f"  {name}: count={len(vals)} min={min(vals):g} max={max(vals):g} "
                f"mean={sum(vals) / len(vals):.4g}"
            )
    return "\n".join(parts), more or bool(approx_rows and approx_rows - 1 > scanned)


def _extract_csv(content: bytes, max_chars: int) -> tuple[str, bool]:
    text = decode_text(content)
    reader = csv.reader(io.StringIO(text))
    summary, more = _summarise_table("CSV", reader, None)
    if not summary:
        raise Rejected("No extractable text")
    return summary[:max_chars], more or len(summary) > max_chars


def _extract_xlsx(content: bytes, max_chars: int) -> tuple[str, bool]:
    _open_zip(content, "Spreadsheet")
    import openpyxl  # deferred: only the worker that needs it pays the import

    workbook = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    try:
        sheets = workbook.worksheets
        blocks: list[str] = []
        truncated = len(sheets) > MAX_SHEETS
        used = 0
        for sheet in sheets[:MAX_SHEETS]:
            approx = sheet.max_row if isinstance(sheet.max_row, int) else None
            rows = sheet.iter_rows(max_row=MAX_SCAN_ROWS + 1, max_col=MAX_COLS, values_only=True)
            summary, more = _summarise_table(f"Sheet '{sheet.title}'", rows, approx)
            truncated = truncated or more
            if summary:
                blocks.append(summary)
                used += len(summary)
            if used >= max_chars:
                truncated = True
                break
        if len(sheets) > MAX_SHEETS:
            blocks.append(f"({len(sheets) - MAX_SHEETS} more sheet(s) not read)")
    finally:
        workbook.close()

    text = "\n\n".join(blocks)
    return text[:max_chars], truncated or len(text) > max_chars


# ── DOCX ────────────────────────────────────────────────────────────────────

_W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def _extract_docx(content: bytes, max_chars: int) -> tuple[str, bool]:
    archive = _open_zip(content, "Word document")
    try:
        member = archive.open("word/document.xml")
    except KeyError:
        raise Rejected("Not a valid Word document (no document body)") from None

    paragraphs: list[str] = []
    length = 0
    truncated = False
    with member:
        for _event, element in ElementTree.iterparse(member, events=("end",)):
            if element.tag == f"{_W}p":
                text = "".join(t.text or "" for t in element.iter(f"{_W}t")).strip()
                if text:
                    paragraphs.append(text)
                    length += len(text) + 1
                element.clear()
                if length >= max_chars:
                    truncated = True
                    break
    text = "\n".join(paragraphs)
    return text[:max_chars], truncated or len(text) > max_chars


# ── PDF ─────────────────────────────────────────────────────────────────────

def _pdf_page_count(pdf) -> int:
    try:
        from pdfminer.pdftypes import resolve1

        return int(resolve1(resolve1(pdf.doc.catalog["Pages"]).get("Count", 0)))
    except Exception:  # noqa: BLE001 — unknown count: assume there may be more
        return MAX_PDF_PAGES + 1


def _extract_pdf(content: bytes, max_chars: int) -> tuple[str, bool]:
    import pdfplumber  # deferred: heavy import

    text_parts: list[str] = []
    length = 0
    truncated = False
    pages = list(range(1, MAX_PDF_PAGES + 1))
    with pdfplumber.open(io.BytesIO(content), pages=pages) as pdf:
        for page in pdf.pages:
            text = page.extract_text()
            if text:
                text_parts.append(text)
                length += len(text)
                if length >= max_chars:
                    truncated = True
                    break
        if len(pdf.pages) >= MAX_PDF_PAGES and _pdf_page_count(pdf) > MAX_PDF_PAGES:
            truncated = True
    text = "\n\n".join(text_parts)
    return text[:max_chars], truncated or len(text) > max_chars
